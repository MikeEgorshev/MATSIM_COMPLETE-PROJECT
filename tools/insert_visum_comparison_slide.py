#!/usr/bin/env python3
"""
Вставить слайд «MATSim vs PTV Visum» в презентацию (после слайда 8).
Использование:
  python tools/insert_visum_comparison_slide.py "C:/Users/.../import.pptx"
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

PARADIGM_TEXT = (
    "Парадигма: это не контраст «микро против макро» по итоговым отчётам — оба подхода выдают "
    "агрегированные показатели (потоки на линках, доли режимов, коридоры). Разница в том, "
    "как задан спрос и согласование: в MATSim — дискретные агенты и планы дня, связка спрос–сеть "
    "через итерации микросимуляции; в Visum — зоны, OD-матрицы и назначение на сеть в логике "
    "классической макромодели (часто 4-step / равновесие) и экосистемы PTV."
)


def patch_visum_slide_paradigm(pptx_path: Path) -> bool:
    """Обновить абзац «Парадигма» на уже вставленном слайде (без дублирования слайда)."""
    prs = Presentation(str(pptx_path))
    changed = False
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            if "Парадигма:" not in sh.text_frame.text:
                continue
            for para in sh.text_frame.paragraphs:
                t = para.text.strip()
                if t.startswith("Парадигма:"):
                    para.text = PARADIGM_TEXT
                    changed = True
                    break
            if changed:
                break
        if changed:
            break
    if changed:
        prs.save(str(pptx_path))
        print(f"OK: paradigm paragraph updated in {pptx_path}")
    else:
        print("No paragraph starting with «Парадигма:» found — nothing changed.")
    return changed


def _move_slide_to_index(prs: Presentation, from_idx: int, to_idx: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    children = list(sld_id_lst)
    el = children[from_idx]
    sld_id_lst.remove(el)
    sld_id_lst.insert(to_idx, el)


def _increment_footers_from_slide(prs: Presentation, start_idx: int) -> None:
    """Увеличить номер в правом нижнем углу на слайдах start_idx..end."""
    for si in range(start_idx, len(prs.slides)):
        slide = prs.slides[si]
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            left_in = sh.left / 914400
            top_in = sh.top / 914400
            if not (11.2 < left_in < 13.2 and 6.5 < top_in < 7.6):
                continue
            t = sh.text_frame.text.strip()
            if t.isdigit():
                sh.text_frame.text = str(int(t) + 1)


def insert_visum_slide(pptx_path: Path) -> None:
    prs = Presentation(str(pptx_path))
    blank = prs.slide_layouts[0]
    new_slide = prs.slides.add_slide(blank)
    _move_slide_to_index(prs, len(prs.slides) - 1, 8)

    slide = prs.slides[8]

    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.16), Inches(12.0), Inches(0.9)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "MATSim и PTV Visum: конкуренция и дополнение"
    p.font.size = Pt(28)
    p.font.bold = True

    body_box = slide.shapes.add_textbox(
        Inches(0.55), Inches(1.12), Inches(12.15), Inches(5.75)
    )
    btf = body_box.text_frame
    btf.word_wrap = True

    lines: list[tuple[str, int, bool]] = [
        (PARADIGM_TEXT, 0, False),
        ("", 0, False),
        ("Конкурируют:", 0, True),
        (
            "выбор «главной» модели города при ограниченном бюджете (лицензии Visum vs открытый MATSim + экспертиза);",
            1,
            False,
        ),
        (
            "оценка загрузки УДС и ОТ для тех же решений — разная методология и входные данные;",
            1,
            False,
        ),
        ("одна команда редко одновременно глубоко владеет обоими стеками.", 1, False),
        ("", 0, False),
        ("Дополняют:", 0, True),
        (
            "Visum — калиброванные OD, экранные линии, стратегические коридоры; MATSim — сценарии «что если», "
            "агенты, согласование спрос–сеть;",
            1,
            False,
        ),
        (
            "пайплайн Шамалгана: зоны/население можно брать из Visum вместо растра — дополнение, не замена;",
            1,
            False,
        ),
        (
            "Берлин: BVG использует PTV Visum; TU Berlin разрабатывает MATSim — сосуществование инструментов в одном городе.",
            1,
            False,
        ),
        ("", 0, False),
        (
            "Схема: Visum (матрицы, зоны, калибровка) → MATSim (агенты, политики, сходимость) → "
            "согласование сценариев / валидация.",
            0,
            False,
        ),
    ]

    first = True
    for text, level, bold in lines:
        if first:
            para = btf.paragraphs[0]
            first = False
        else:
            para = btf.add_paragraph()
        para.text = text
        para.level = level
        para.font.size = Pt(13)
        para.font.bold = bold
        para.space_after = Pt(3)

    foot = slide.shapes.add_textbox(
        Inches(12.4), Inches(7.05), Inches(0.75), Inches(0.4)
    )
    fp = foot.text_frame.paragraphs[0]
    fp.text = "9"
    fp.font.size = Pt(12)

    _increment_footers_from_slide(prs, 9)

    prs.save(str(pptx_path))
    print(f"OK: inserted slide at position 9, updated footers 9→10… in {pptx_path}")


def main() -> None:
    if sys.platform == "win32":
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path, help="Путь к .pptx")
    ap.add_argument(
        "--patch-paradigm-only",
        action="store_true",
        help="Только заменить абзац «Парадигма» на слайде Visum (слайд уже должен быть в презентации).",
    )
    args = ap.parse_args()
    if not args.pptx.is_file():
        raise SystemExit(f"Not found: {args.pptx}")
    path = args.pptx.resolve()
    if args.patch_paradigm_only:
        patch_visum_slide_paradigm(path)
    else:
        insert_visum_slide(path)


if __name__ == "__main__":
    main()
