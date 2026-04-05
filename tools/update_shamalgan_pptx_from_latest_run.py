#!/usr/bin/env python3
"""
Обновить метрики последнего расчёта в пользовательском .pptx (не пересобирая из git).

По умолчанию:
- правит таблицу «Доля авто / ОТ / пеших» и «Средняя полезность» на слайде с результатами;
- обновляет поясняющий абзац про ОТ;
- подменяет встроенные PNG графика сходимости (слайды с картинкой справа сверху, как в типовой вёрстке).

Использование:
  python tools/update_shamalgan_pptx_from_latest_run.py "C:/Users/.../Downloads/10_final_report_shamalgan.pptx"
"""

from __future__ import annotations

import argparse
import re
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONVERGENCE_PNG = ROOT / "Visualization" / "docs-progress" / "simulation_convergence.png"

# Последний расчёт (output_100it): см. modestats.csv / scorestats.csv / traffic_stats
MODE_CAR = "56,0%"
MODE_PT = "19,3%"
MODE_WALK = "24,7%"
UTIL_ROW = "124,1 (на ит. 0: −24,0)"
FOOTER_OT = (
    "Доля поездок на ОТ выросла с ~12,5% (ит. 0) до 19,3% (ит. 100); доля автомобиля остаётся "
    "доминирующей, но ниже, чем в ранних черновых отчётах с другими настройками сценария — "
    "это следует интерпретировать в связке с текущей сетью, расписанием ОТ и штрафами режимов в config.xml."
)
SPEED_CELL = "23,0 км/ч"


def update_presentation(path: Path, png: Path) -> None:
    blob = png.read_bytes()
    prs = Presentation(str(path))

    # Слайд «Запуск и результаты» — ищем по заголовку
    for slide in prs.slides:
        title = ""
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                t0 = sh.text_frame.text.strip().split("\n")[0]
                if "Запуск и результаты" in t0 or "результаты симуляции" in t0.lower():
                    title = t0
                    break
        if not title:
            continue

        cell_map = [
            ("126,3 (рост с 104,9)", UTIL_ROW),
            ("79,2%", MODE_CAR),
            ("79.2%", "56.0%"),
            ("4,9%", MODE_PT),
            ("4.9%", "19.3%"),
            ("15,9%", MODE_WALK),
            ("15.9%", "24.7%"),
            ("126,3", "124,1"),
            ("рост с 104,9", "на ит. 0: −24,0"),
        ]

        for sh in slide.shapes:
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        txt = cell.text
                        new = txt
                        for old, rep in cell_map:
                            if old in new:
                                new = new.replace(old, rep)
                        if re.fullmatch(r"23,\d+\s*км/ч", new.strip()):
                            new = SPEED_CELL
                        if new != txt:
                            cell.text = new
            if sh.has_text_frame:
                ptxt = sh.text_frame.text
                if "Доля ОТ" in ptxt and "снизилась" in ptxt:
                    sh.text_frame.text = FOOTER_OT
                elif "16,5%" in ptxt or ("4,9%" in ptxt and "снизилась" in ptxt):
                    sh.text_frame.text = FOOTER_OT

        # Картинка сходимости — верхняя правая (top < 2 in)
        for sh in list(slide.shapes):
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            if sh.left / 914400 >= 5.5 and sh.top / 914400 < 2.0:
                left, top, width, height = sh.left, sh.top, sh.width, sh.height
                el = sh._element
                parent = el.getparent()
                parent.remove(el)
                slide.shapes.add_picture(BytesIO(blob), left, top, width=width, height=height)
                break

    # Слайд «Оценка плана, итерации…» — обновить график
    for slide in prs.slides:
        hit = False
        for sh in slide.shapes:
            if sh.has_text_frame and "Оценка плана" in sh.text_frame.text:
                hit = True
                break
        if not hit:
            continue
        for sh in list(slide.shapes):
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            if sh.left / 914400 >= 5.5 and sh.top / 914400 < 2.0:
                left, top, width, height = sh.left, sh.top, sh.width, sh.height
                el = sh._element
                parent = el.getparent()
                parent.remove(el)
                slide.shapes.add_picture(BytesIO(blob), left, top, width=width, height=height)
                break

    prs.save(str(path))
    print(f"Updated: {path}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", type=Path, help="Путь к .pptx")
    ap.add_argument("--png", type=Path, default=DEFAULT_CONVERGENCE_PNG, help="Новый simulation_convergence.png")
    args = ap.parse_args()
    if not args.pptx.is_file():
        raise SystemExit(f"Не найден файл: {args.pptx}")
    if not args.png.is_file():
        raise SystemExit(f"Не найден PNG: {args.png}")
    update_presentation(args.pptx.resolve(), args.png.resolve())


if __name__ == "__main__":
    main()
